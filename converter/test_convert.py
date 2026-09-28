"""Converter unit tests — real Pandoc on a generated DOCX, real pymupdf4llm on a generated PDF.

Runs inside the converter image (pandoc + python-docx + pymupdf available). Verifies the
pure bytes->markdown contract: GFM tables, data-URI images, no on-disk media
leakage. LaTeX formulas (OMML) are not exercised here — python-docx cannot emit
OMML; formula round-trip is covered by manual testing with a real Word file.
"""

import io
import struct
import zlib

import pytest

from convert import docx_bytes_to_markdown, pdf_bytes_to_markdown


def _valid_png(w: int, h: int) -> bytes:
    """Build a structurally valid RGB PNG (correct chunk CRCs) of any size."""
    def chunk(tag: bytes, data: bytes) -> bytes:
        return (struct.pack(">I", len(data)) + tag + data
                + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF))

    sig = b"\x89PNG\r\n\x1a\n"
    ihdr = struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)  # 8-bit, RGB
    raw = b"".join(b"\x00" + b"\x10\x20\x30" * w for _ in range(h))  # filtered scanlines
    idat = zlib.compress(raw)
    return sig + chunk(b"IHDR", ihdr) + chunk(b"IDAT", idat) + chunk(b"IEND", b"")


def _valid_png_1x1() -> bytes:
    return _valid_png(1, 1)


def _make_docx() -> bytes:
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    doc.add_heading("Title", level=1)
    doc.add_paragraph("Hello world.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "a"
    table.cell(0, 1).text = "b"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"

    doc.add_picture(io.BytesIO(_valid_png_1x1()), width=Inches(1.0))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_to_markdown_has_heading_and_table_and_inline_image():
    md = docx_bytes_to_markdown(_make_docx())
    assert "Title" in md
    # GFM table delimiter row.
    assert "|" in md and "---" in md
    # Image inlined as a base64 data URI, not an on-disk path.
    assert "data:image/" in md and "base64," in md
    assert "media/" not in md


def test_markdown_to_docx_produces_bytes():
    from convert import markdown_to_docx
    data = markdown_to_docx("# Hello\n\nSome **bold** text.")
    assert len(data) > 100
    assert data[:4] == b"PK\x03\x04"


# --- PDF fixtures (built by PyMuPDF itself) -----------------------------------


def _make_text_table_pdf() -> bytes:
    """One page: a big-font title, a body line, and a drawn 2x2 grid table."""
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 90), "Big Title", fontsize=24, fontname="helv")
    page.insert_text((72, 120), "Hello PDF world.", fontsize=11, fontname="helv")
    x0, y0, cw, ch = 72, 160, 100, 24
    cells = (("a", "b"), ("1", "2"))
    shape = page.new_shape()
    for r in range(3):
        shape.draw_line(pymupdf.Point(x0, y0 + r * ch), pymupdf.Point(x0 + 2 * cw, y0 + r * ch))
    for c in range(3):
        shape.draw_line(pymupdf.Point(x0 + c * cw, y0), pymupdf.Point(x0 + c * cw, y0 + 2 * ch))
    shape.finish(color=(0, 0, 0), width=1)
    shape.commit()
    for r, row in enumerate(cells):
        for c, val in enumerate(row):
            page.insert_text(pymupdf.Point(x0 + c * cw + 6, y0 + r * ch + 16), val,
                             fontsize=10, fontname="helv")
    return doc.tobytes()


def _make_images_pdf() -> bytes:
    """Page 1: a 20x20 image (below the min-side filter). Page 2: a 300x300 one."""
    import pymupdf

    doc = pymupdf.open()
    p1 = doc.new_page()
    p1.insert_image(pymupdf.Rect(72, 120, 92, 140), stream=_valid_png(20, 20))
    p2 = doc.new_page()
    p2.insert_image(pymupdf.Rect(72, 120, 272, 320), stream=_valid_png(300, 300))
    return doc.tobytes()


def _make_repeat_image_pdf(pages: int = 3) -> bytes:
    """The same 300x300 image on N consecutive pages."""
    import pymupdf

    img = _valid_png(300, 300)
    doc = pymupdf.open()
    for _ in range(pages):
        page = doc.new_page()
        page.insert_image(pymupdf.Rect(72, 120, 272, 320), stream=img)
    return doc.tobytes()


def test_pdf_to_markdown_has_text_and_table():
    md = pdf_bytes_to_markdown(_make_text_table_pdf())
    assert "Big Title" in md
    assert "Hello PDF world." in md
    # GFM table delimiter row.
    assert "|" in md and "---" in md
    assert "a" in md and "1" in md


def test_pdf_images_min_side_filter_and_labelled_definitions():
    md = pdf_bytes_to_markdown(_make_images_pdf())
    # Only the 300x300 image survives; it appears ONCE as a labelled definition.
    assert md.count("data:image/") == 1
    # Definition + usage forms the backend's _REF_DEF_PATTERN collapse.
    assert "![](data:" not in md  # never the inline form the DOCX engine emits
    import re
    assert re.search(r"^!\[image\]\[img-[0-9a-f]+\]$", md, re.MULTILINE)
    assert re.search(r"^\[img-[0-9a-f]+\]: data:image/(png|jpeg);base64,", md, re.MULTILINE)


def test_pdf_duplicate_image_one_definition_many_usages():
    md = pdf_bytes_to_markdown(_make_repeat_image_pdf(pages=3))
    assert md.count("data:image/") == 1
    assert md.count("![image][img-") == 3


def test_pdf_garbage_bytes_raise():
    import pymupdf
    with pytest.raises(pymupdf.FileDataError):
        pdf_bytes_to_markdown(b"this is not a pdf document")
